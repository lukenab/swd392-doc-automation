from __future__ import annotations

from datetime import date
import json
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .loader import ProjectData


FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
TABLE_CONTENT_FONT_SIZE = Pt(11)
TABLE_LABEL_FONT_SIZE = Pt(12)
USE_CASE_HEADING_FONT_SIZE = Pt(13)
DIAGRAM_MAX_WIDTH = Cm(16.5)
DIAGRAM_MAX_HEIGHT = Cm(20.5)
TABLE_HEADER_FILL = "1F4E78"
TABLE_HEADER_TEXT = RGBColor(255, 255, 255)


def _format_date(value: Any) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    text = str(value or "")
    try:
        return date.fromisoformat(text).strftime("%d/%m/%Y")
    except ValueError:
        return text


def _escape_markdown(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", "<br>")


def _actor_text(use_case: dict[str, Any]) -> str:
    actors = [use_case["primary_actor"], *(use_case.get("secondary_actors") or [])]
    return ", ".join(dict.fromkeys(actors))


def _use_case_id(use_case: dict[str, Any]) -> str:
    return str(use_case["_display_id"])


def _group_name(project: ProjectData, item: dict[str, Any]) -> str:
    return str(project.groups[item["group"]]["name"])


def _business_rule_ids(project: ProjectData, use_case: dict[str, Any]) -> str:
    keys = use_case.get("business_rules", []) or []
    return ", ".join(project.business_rules[key]["_display_id"] for key in keys) or "N/A"


def _business_rule_name(rule: dict[str, Any]) -> str:
    semantic_name = str(rule["key"]).removeprefix("BR-")
    expanded_tokens = {
        "PO": "Product Owner",
        "SYSADMIN": "System Administrator",
    }
    return " ".join(
        expanded_tokens.get(token, token.title())
        for token in semantic_name.split("-")
    )


def _numbered(items: Iterable[Any]) -> str:
    values = list(items or [])
    if not values:
        return "N/A"
    return "\n".join(f"{index}. {item}" for index, item in enumerate(values, start=1))


def _normal_flow_text(use_case: dict[str, Any]) -> str:
    return "\n".join(
        f"{index}. {step['actor']}: {step['action']}"
        for index, step in enumerate(use_case.get("normal_flow", []), start=1)
    ) or "N/A"


def _alternative_flow_text(use_case: dict[str, Any]) -> str:
    flows = use_case.get("alternative_flows", []) or []
    if not flows:
        return "N/A"
    parts: list[str] = []
    for index, flow in enumerate(flows, start=1):
        if isinstance(flow, str):
            parts.append(f"AF-{index:02d}: {flow}")
            continue
        heading = f"{flow.get('id', 'AF')}: {flow.get('condition', '')}".strip()
        steps = flow.get("steps", []) or []
        parts.append(heading + ("\n" + _numbered(steps) if steps else ""))
    return "\n\n".join(parts)


def _exception_text(use_case: dict[str, Any]) -> str:
    exceptions = use_case.get("exceptions", []) or []
    if not exceptions:
        return "N/A"
    lines: list[str] = []
    for index, item in enumerate(exceptions, start=1):
        if isinstance(item, str):
            lines.append(f"EX-{index:02d}: {item}")
        else:
            lines.append(f"{item.get('id', 'EX')}: {item.get('description', '')}")
    return "\n".join(lines)


def _detail_fields(project: ProjectData, use_case: dict[str, Any]) -> list[tuple[str, str]]:
    return [
        ("Trigger", str(use_case.get("trigger") or "N/A")),
        ("Description", str(use_case.get("description") or "N/A")),
        ("Preconditions", _numbered(use_case.get("preconditions", []))),
        ("Postconditions", _numbered(use_case.get("postconditions", []))),
        ("Normal Flow", _normal_flow_text(use_case)),
        ("Alternative Flows", _alternative_flow_text(use_case)),
        ("Exceptions", _exception_text(use_case)),
        ("Priority", str(use_case.get("priority") or "N/A")),
        ("Frequency of Use", str(use_case.get("frequency_of_use") or "N/A")),
        ("Business Rules", _business_rule_ids(project, use_case)),
        ("Other Information", str(use_case.get("other_information") or "N/A")),
        ("Assumptions", _numbered(use_case.get("assumptions", []))),
    ]


def _generate_markdown(project: ProjectData, use_cases: list[dict[str, Any]], output_dir: Path) -> list[Path]:
    summary_lines = [
        "# Use Case List",
        "",
        "| ID | Use Case | Actors | Use Case Description |",
        "|---|---|---|---|",
    ]
    for use_case in use_cases:
        summary_lines.append(
            "| {id} | {name} | {actors} | {summary} |".format(
                id=_escape_markdown(_use_case_id(use_case)),
                name=_escape_markdown(use_case["name"]),
                actors=_escape_markdown(_actor_text(use_case)),
                summary=_escape_markdown(use_case["summary"]),
            )
        )

    details_lines = ["# Use Case Descriptions", ""]
    current_group = None
    for use_case in use_cases:
        if use_case["group"] != current_group:
            current_group = use_case["group"]
            details_lines.extend([f"## {_group_name(project, use_case)}", ""])
        details_lines.extend(
            [
                f"### {_use_case_id(use_case)} - {use_case['name']}",
                "",
                "| Field | Value |",
                "|---|---|",
                f"| UC ID and Name | {_escape_markdown(_use_case_id(use_case))} - {_escape_markdown(use_case['name'])} |",
                f"| Created By | {_escape_markdown(use_case['created_by'])} |",
                f"| Date Created | {_escape_markdown(_format_date(use_case['date_created']))} |",
                f"| Primary Actor | {_escape_markdown(use_case['primary_actor'])} |",
                f"| Secondary Actors | {_escape_markdown(', '.join(use_case.get('secondary_actors', [])) or 'None')} |",
            ]
        )
        for label, value in _detail_fields(project, use_case):
            details_lines.append(f"| {_escape_markdown(label)} | {_escape_markdown(value)} |")
        details_lines.append("")

    summary_path = output_dir / "use-case-list.md"
    details_path = output_dir / "use-case-descriptions.md"
    mapping_path = output_dir / "id-mapping.md"
    summary_path.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    details_path.write_text("\n".join(details_lines) + "\n", encoding="utf-8")

    mapping_lines = [
        "# Generated ID Mapping",
        "",
        "## Use Cases",
        "",
        "| Generated ID | Semantic Key | Domain | Name | Order |",
        "|---|---|---|---|---|",
    ]
    for use_case in use_cases:
        mapping_lines.append(
            f"| {_use_case_id(use_case)} | `{use_case['key']}` | "
            f"{_escape_markdown(_group_name(project, use_case))} | "
            f"{_escape_markdown(use_case['name'])} | {use_case['order']} |"
        )
    mapping_lines.extend(
        [
            "",
            "## Business Rules",
            "",
            "| Generated ID | Semantic Key | Domain | Description | Order |",
            "|---|---|---|---|---|",
        ]
    )
    for rule in project.business_rules.values():
        mapping_lines.append(
            f"| {rule['_display_id']} | `{rule['key']}` | "
            f"{_escape_markdown(_group_name(project, rule))} | "
            f"{_escape_markdown(rule['description'])} | {rule['order']} |"
        )
    mapping_path.write_text("\n".join(mapping_lines) + "\n", encoding="utf-8")
    business_rule_path = _generate_business_rules_markdown(project, output_dir)
    return [summary_path, details_path, business_rule_path, mapping_path]


def _generate_business_rules_markdown(project: ProjectData, output_dir: Path) -> Path:
    lines = [
        "# Business Rules",
        "",
        "This section defines the business constraints that govern system behavior.",
        "",
        "| ID | Business Rule | Description |",
        "|---|---|---|",
    ]
    for rule in project.business_rules.values():
        lines.append(
            f"| {rule['_display_id']} | {_escape_markdown(_business_rule_name(rule))} | "
            f"{_escape_markdown(rule['description'])} |"
        )
    lines.append("")

    output_path = output_dir / "business-rule-list.md"
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def _set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_run_font(run, size=FONT_SIZE, color: RGBColor | None = None) -> None:
    run.font.name = FONT_NAME
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _set_cell_margins(cell) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for side, value in {"top": 90, "left": 100, "bottom": 90, "right": 100}.items():
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _set_cell_text(
    cell,
    text: str,
    bold: bool = False,
    center: bool = False,
    color: RGBColor | None = None,
    size=FONT_SIZE,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    for line_number, line in enumerate(str(text).split("\n")):
        if line_number:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        _set_run_font(run, size=size, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        row.height = Cm(0.75)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _set_table_column_widths(table, widths: list[Cm]) -> None:
    for index, width in enumerate(widths):
        table.columns[index].width = width
        for cell in table.columns[index].cells:
            cell.width = width


def _enforce_document_font(document: Document) -> None:
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = (
            USE_CASE_HEADING_FONT_SIZE if style_name == "Heading 2" else FONT_SIZE
        )
        style.font.color.rgb = RGBColor(0, 0, 0)
        style_properties = style._element.get_or_add_rPr()
        fonts = style_properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style_properties.insert(0, fonts)
        for attribute in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    for paragraph in document.paragraphs:
        font_size = (
            USE_CASE_HEADING_FONT_SIZE
            if paragraph.style.name == "Heading 2"
            else FONT_SIZE
        )
        for run in paragraph.runs:
            _set_run_font(run, font_size)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_font(run, size=None)


def _load_diagram_images(diagram_dir: Path | None) -> list[tuple[str, Path]]:
    if diagram_dir is None:
        return []

    root = diagram_dir.resolve()
    if not root.is_dir():
        raise ValueError(f"Diagram directory does not exist: {root}")

    manifest_path = root / "manifest.json"
    images: list[tuple[str, Path]] = []
    if manifest_path.is_file():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            raise ValueError(f"Cannot read diagram manifest: {manifest_path}: {error}") from error

        for item in manifest.get("diagrams", []):
            relative_path = item.get("files", {}).get("png")
            if not relative_path:
                continue
            image_path = (root / relative_path).resolve()
            if root != image_path and root not in image_path.parents:
                raise ValueError(f"Diagram path escapes its output directory: {relative_path}")
            if not image_path.is_file():
                raise ValueError(f"Diagram image listed in manifest does not exist: {image_path}")
            images.append((str(item.get("title") or item.get("key") or image_path.stem), image_path))
    else:
        images = [(path.stem, path) for path in sorted(root.glob("use-case-*.png"))]

    if not images:
        raise ValueError(f"No PNG use case diagram found in: {root}")
    return images


def _add_diagram_section(document: Document, images: list[tuple[str, Path]]) -> None:
    document.add_heading("II.5.2.1 Use Case Diagram", level=1)
    for index, (title, image_path) in enumerate(images, start=1):
        image = Image.from_file(str(image_path))
        scale = min(
            int(DIAGRAM_MAX_WIDTH) / int(image.width),
            int(DIAGRAM_MAX_HEIGHT) / int(image.height),
        )
        width = int(int(image.width) * scale)
        height = int(int(image.height) * scale)

        picture_paragraph = document.add_paragraph()
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.keep_with_next = True
        shape = picture_paragraph.add_run().add_picture(
            str(image_path),
            width=width,
            height=height,
        )
        shape._inline.docPr.set("descr", title)

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(6)
        caption_run = caption.add_run(f"Figure II.5.2.1-{index}. {title}")
        caption_run.italic = True
        _set_run_font(caption_run)

        if index < len(images):
            document.add_page_break()


def _generate_docx(
    project: ProjectData,
    use_cases: list[dict[str, Any]],
    output_dir: Path,
    diagram_dir: Path | None = None,
) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    diagram_images = _load_diagram_images(diagram_dir)
    if diagram_images:
        _add_diagram_section(document, diagram_images)
        document.add_page_break()

    document.add_heading("II.5.2.2 Use Case Descriptions", level=1)

    summary_table = document.add_table(rows=1, cols=4)
    summary_table.style = "Table Grid"
    _style_table(summary_table)
    headers = ["ID", "Use Case", "Actors", "Use Case Description"]
    for index, header in enumerate(headers):
        _set_cell_text(
            summary_table.rows[0].cells[index],
            header,
            bold=True,
            center=True,
            color=TABLE_HEADER_TEXT,
            size=TABLE_LABEL_FONT_SIZE,
        )
        _set_cell_shading(summary_table.rows[0].cells[index], TABLE_HEADER_FILL)
    widths = [Cm(1.6), Cm(3.6), Cm(3.7), Cm(8.1)]
    for use_case in use_cases:
        cells = summary_table.add_row().cells
        values = [_use_case_id(use_case), use_case["name"], _actor_text(use_case), use_case["summary"]]
        for index, value in enumerate(values):
            _set_cell_text(
                cells[index],
                value,
                center=index == 0,
                size=TABLE_CONTENT_FONT_SIZE,
            )
            cells[index].width = widths[index]
    _set_table_column_widths(summary_table, widths)
    _style_table(summary_table)

    for use_case in use_cases:
        heading = document.add_heading(f"{_use_case_id(use_case)} - {use_case['name']}", level=2)
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(6)
        heading.paragraph_format.keep_with_next = True

        table = document.add_table(rows=3, cols=4)
        table.style = "Table Grid"
        _style_table(table)

        first_row = table.rows[0].cells
        _set_cell_text(
            first_row[0], "UC ID and Name", bold=True, size=TABLE_LABEL_FONT_SIZE
        )
        merged = first_row[1].merge(first_row[3])
        _set_cell_text(
            merged,
            f"{_use_case_id(use_case)} - {use_case['name']}",
            size=TABLE_CONTENT_FONT_SIZE,
        )

        second_row = table.rows[1].cells
        _set_cell_text(second_row[0], "Created By", bold=True, size=TABLE_LABEL_FONT_SIZE)
        _set_cell_text(second_row[1], use_case["created_by"], size=TABLE_CONTENT_FONT_SIZE)
        _set_cell_text(second_row[2], "Date Created", bold=True, size=TABLE_LABEL_FONT_SIZE)
        _set_cell_text(
            second_row[3],
            _format_date(use_case["date_created"]),
            size=TABLE_CONTENT_FONT_SIZE,
        )

        third_row = table.rows[2].cells
        _set_cell_text(third_row[0], "Primary Actor", bold=True, size=TABLE_LABEL_FONT_SIZE)
        _set_cell_text(third_row[1], use_case["primary_actor"], size=TABLE_CONTENT_FONT_SIZE)
        _set_cell_text(
            third_row[2], "Secondary Actors", bold=True, size=TABLE_LABEL_FONT_SIZE
        )
        _set_cell_text(
            third_row[3],
            ", ".join(use_case.get("secondary_actors", [])) or "None",
            size=TABLE_CONTENT_FONT_SIZE,
        )

        for label, value in _detail_fields(project, use_case):
            cells = table.add_row().cells
            _set_cell_text(cells[0], label, bold=True, size=TABLE_LABEL_FONT_SIZE)
            merged_value = cells[1].merge(cells[3])
            _set_cell_text(merged_value, value, size=TABLE_CONTENT_FONT_SIZE)

        for row in table.rows:
            row.cells[0].width = Cm(3.7)
        _style_table(table)

    output_path = output_dir / "use-case-descriptions.docx"
    _enforce_document_font(document)
    document.save(output_path)
    return output_path


def _generate_business_rules_docx(project: ProjectData, output_dir: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    heading = document.add_heading("II.4 Business Rules", level=1)
    heading.paragraph_format.space_after = Pt(6)

    introduction = document.add_paragraph(
        "This section defines the business constraints that govern system behavior."
    )
    introduction.paragraph_format.space_after = Pt(8)

    widths = [Cm(1.5), Cm(5.2), Cm(10.3)]
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["ID", "Business Rule", "Description"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_text(
            cell,
            header,
            bold=True,
            center=True,
            color=TABLE_HEADER_TEXT,
        )
        _set_cell_shading(cell, TABLE_HEADER_FILL)
        cell.width = widths[index]
    for rule in project.business_rules.values():
        cells = table.add_row().cells
        values = [
            rule["_display_id"],
            _business_rule_name(rule),
            rule["description"],
        ]
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, center=index == 0)
            cells[index].width = widths[index]

    _set_table_column_widths(table, widths)
    _style_table(table)

    output_path = output_dir / "business-rules.docx"
    _enforce_document_font(document)
    document.save(output_path)
    return output_path


def _plantuml_alias(value: str) -> str:
    return "".join(character if character.isalnum() else "_" for character in value)


def _generate_plantuml(project: ProjectData, use_cases: list[dict[str, Any]], output_dir: Path) -> Path:
    used_actor_names = {
        actor
        for use_case in use_cases
        for actor in [use_case["primary_actor"], *(use_case.get("secondary_actors") or [])]
    }
    lines = [
        "@startuml",
        "left to right direction",
        "skinparam packageStyle rectangle",
        "skinparam actorStyle awesome",
        "",
    ]
    for actor_name in sorted(used_actor_names):
        lines.append(f'actor "{actor_name}" as ACT_{_plantuml_alias(actor_name)}')
    lines.extend(["", 'rectangle "Project Management System" {'])
    current_group = None
    for use_case in use_cases:
        if use_case["group"] != current_group:
            if current_group is not None:
                lines.append("  }")
            current_group = use_case["group"]
            lines.append(
                f'  package "{_group_name(project, use_case)}" '
                f'as DOMAIN_{_plantuml_alias(current_group)} {{'
            )
        lines.append(
            f'    usecase "{_use_case_id(use_case)}\\n{use_case["name"]}" '
            f'as {_plantuml_alias(use_case["key"])}'
        )
    if current_group is not None:
        lines.append("  }")
    lines.append("}")
    lines.append("")
    for use_case in use_cases:
        use_case_alias = _plantuml_alias(use_case["key"])
        actors = [use_case["primary_actor"], *(use_case.get("secondary_actors") or [])]
        for actor_name in dict.fromkeys(actors):
            lines.append(f"ACT_{_plantuml_alias(actor_name)} --> {use_case_alias}")
    lines.extend(["", "@enduml", ""])

    output_path = output_dir / "use-case-diagram.puml"
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def build_outputs(
    project: ProjectData,
    use_cases: list[dict[str, Any]],
    output_dir: Path,
    output_format: str,
    diagram_dir: Path | None = None,
) -> list[Path]:
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []

    if output_format in {"all", "markdown"}:
        generated.extend(_generate_markdown(project, use_cases, output_dir))
    if output_format in {"all", "docx"}:
        generated.append(_generate_docx(project, use_cases, output_dir, diagram_dir))
        generated.append(_generate_business_rules_docx(project, output_dir))
    if output_format == "business-rules":
        generated.append(_generate_business_rules_markdown(project, output_dir))
        generated.append(_generate_business_rules_docx(project, output_dir))
    if output_format in {"all", "plantuml"}:
        generated.append(_generate_plantuml(project, use_cases, output_dir))
    return generated
