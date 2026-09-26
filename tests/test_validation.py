from copy import deepcopy
import base64
import json
from pathlib import Path
import shutil
import unittest

import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from usecase_tool.generator import (
    _alternative_flow_text,
    _exception_text,
    _generate_business_rules_docx,
    _generate_business_rules_markdown,
    _generate_docx,
)
from usecase_tool.loader import assign_display_ids, load_project
from usecase_tool.validator import validate_project


ROOT = Path(__file__).resolve().parents[1]


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.project = load_project(ROOT / "data")

    def test_sample_data_is_valid(self):
        self.assertEqual([], validate_project(self.project))

    def test_project_roles_and_scrum_accountabilities_are_modeled_as_user(self):
        merged_role_names = {
            "Project Member",
            "Project Owner",
            "Product Owner",
            "Developer",
        }
        self.assertEqual(
            {
                "Visitor",
                "System Administrator",
                "User",
                "Email Service",
                "Identity Provider",
            },
            set(self.project.actors),
        )

        project_use_cases = [
            item
            for item in self.project.use_cases
            if item["group"] not in {"ACCOUNT_AUTHENTICATION", "SYSTEM_ADMINISTRATION"}
        ]
        self.assertTrue(project_use_cases)
        for use_case in project_use_cases:
            self.assertEqual("User", use_case["primary_actor"])
            self.assertTrue(
                merged_role_names.isdisjoint(use_case.get("secondary_actors", []))
            )
            self.assertTrue(
                all(
                    step["actor"] not in merged_role_names
                    for step in use_case.get("normal_flow", [])
                )
            )

    def test_email_notifications_are_in_iteration_scope(self):
        feature_data = yaml.safe_load((ROOT / "data" / "major_features.yml").read_text(encoding="utf-8"))
        feature = next(item for item in feature_data["major_features"] if item["key"] == "FE-08")
        notification_use_case = next(
            item for item in self.project.use_cases if item["key"] == "UC-NOTIFICATIONS-REVIEW"
        )
        comment_use_case = next(
            item for item in self.project.use_cases if item["key"] == "UC-WORK-ITEM-COMMENT"
        )

        self.assertIn("email notifications", feature["description"])
        self.assertIn("Email Service", comment_use_case["secondary_actors"])
        self.assertIn("email delivery", notification_use_case["other_information"])
        self.assertIn("push delivery remains outside", notification_use_case["other_information"].lower())
        self.assertEqual(
            [
                "BR-NOTIFICATION-OWNER-ONLY",
                "BR-PROJECT-MEMBER-ACCESS",
                "BR-NOTIFICATION-EMAIL-PREFERENCE",
            ],
            notification_use_case["business_rules"],
        )
        self.assertTrue(
            any(
                "email notification cannot be delivered" in exception["description"]
                for exception in comment_use_case["exceptions"]
            )
        )

    def test_authentication_use_cases_and_rules_are_synchronized(self):
        authentication_cases = [
            item
            for item in self.project.use_cases
            if item["group"] == "ACCOUNT_AUTHENTICATION"
        ]
        self.assertEqual(
            ["UC-01", "UC-02", "UC-03", "UC-04"],
            [item["_display_id"] for item in authentication_cases],
        )
        self.assertEqual(
            ["Register Account", "Sign In", "Sign Out", "Reset Password"],
            [item["name"] for item in authentication_cases],
        )
        self.assertEqual("Visitor", authentication_cases[0]["primary_actor"])
        self.assertIn("Identity Provider", authentication_cases[0]["secondary_actors"])
        self.assertIn("Email Service", authentication_cases[3]["secondary_actors"])

        account_rule = self.project.business_rules["BR-SYSADMIN-USER-ACCOUNT"]
        self.assertNotIn("create", account_rule["description"].lower())
        for use_case in authentication_cases:
            for rule_key in use_case["business_rules"]:
                self.assertIn(rule_key, self.project.business_rules)

    def test_audit_feedback_changes_are_synchronized(self):
        use_cases = {item["key"]: item for item in self.project.use_cases}

        self.assertIn(
            "BR-PROJECT-OWNER-ADMINISTRATION",
            use_cases["UC-PROJECT-OWNERSHIP-TRANSFER"]["business_rules"],
        )
        self.assertNotIn(
            "BR-PRODUCT-OWNER-BACKLOG",
            use_cases["UC-PRODUCT-GOAL-MANAGE"]["business_rules"],
        )
        self.assertIn(
            "BR-SPRINT-ONE-DRAFT",
            use_cases["UC-SPRINT-PLAN"]["business_rules"],
        )
        self.assertEqual(
            ["BR-NOTIFICATION-OWNER-ONLY", "BR-PROJECT-MEMBER-ACCESS", "BR-NOTIFICATION-EMAIL-PREFERENCE"],
            use_cases["UC-NOTIFICATIONS-REVIEW"]["business_rules"],
        )
        self.assertIn(
            "BR-AUDIT-VIEW-SYSADMIN-ONLY",
            use_cases["UC-SYSTEM-AUDIT-LOG-REVIEW"]["business_rules"],
        )
        self.assertTrue(
            any(
                exception["id"] == "EX-04"
                for exception in use_cases["UC-SIGN-IN"]["exceptions"]
            )
        )
        self.assertEqual(
            "BR-14",
            self.project.business_rules["BR-SPRINT-GOAL-REQUIRED"]["_display_id"],
        )

    def test_unknown_actor_is_reported(self):
        changed = deepcopy(self.project.use_cases)
        changed[0]["primary_actor"] = "Unknown Actor"
        invalid_project = self.project.__class__(
            data_dir=self.project.data_dir,
            groups=self.project.groups,
            actors=self.project.actors,
            business_rules=self.project.business_rules,
            use_cases=changed,
        )
        errors = validate_project(invalid_project)
        self.assertTrue(any("unknown primary actor" in error for error in errors))

    def test_unknown_business_rule_is_reported(self):
        changed = deepcopy(self.project.use_cases)
        changed[0]["business_rules"].append("BR-UNKNOWN-RULE")
        invalid_project = self.project.__class__(
            data_dir=self.project.data_dir,
            groups=self.project.groups,
            actors=self.project.actors,
            business_rules=self.project.business_rules,
            use_cases=changed,
        )
        errors = validate_project(invalid_project)
        self.assertTrue(any("BR-UNKNOWN-RULE" in error for error in errors))

    def test_display_ids_are_generated_from_domain_and_order(self):
        groups = {
            "PROJECT": {"key": "PROJECT", "order": 100},
            "TASK": {"key": "TASK", "order": 400},
        }
        use_cases = [
            {"key": "UC-TASK-CREATE", "group": "TASK", "order": 100},
            {"key": "UC-PROJECT-ARCHIVE", "group": "PROJECT", "order": 200},
            {"key": "UC-PROJECT-CREATE", "group": "PROJECT", "order": 100},
        ]
        assign_display_ids(use_cases, "UC", groups)
        self.assertEqual(
            [
                ("UC-PROJECT-CREATE", "UC-01"),
                ("UC-PROJECT-ARCHIVE", "UC-02"),
                ("UC-TASK-CREATE", "UC-03"),
            ],
            [(item["key"], item["_display_id"]) for item in use_cases],
        )

    def test_use_case_domain_order_can_differ_from_business_rule_order(self):
        groups = {
            "PROJECT": {"key": "PROJECT", "order": 100, "use_case_order": 200},
            "AUTH": {"key": "AUTH", "order": 200, "use_case_order": 100},
        }
        use_cases = [
            {"key": "UC-PROJECT", "group": "PROJECT", "order": 100},
            {"key": "UC-SIGN-IN", "group": "AUTH", "order": 100},
        ]
        rules = [
            {"key": "BR-PROJECT", "group": "PROJECT", "order": 100},
            {"key": "BR-AUTH", "group": "AUTH", "order": 100},
        ]

        assign_display_ids(use_cases, "UC", groups, group_order_field="use_case_order")
        assign_display_ids(rules, "BR", groups)

        self.assertEqual("UC-SIGN-IN", use_cases[0]["key"])
        self.assertEqual("BR-PROJECT", rules[0]["key"])

    def test_non_integer_order_is_reported(self):
        changed = deepcopy(self.project.use_cases)
        changed[0]["order"] = "first"
        invalid_project = self.project.__class__(
            data_dir=self.project.data_dir,
            groups=self.project.groups,
            actors=self.project.actors,
            business_rules=self.project.business_rules,
            use_cases=changed,
        )
        errors = validate_project(invalid_project)
        self.assertTrue(any("order must be an integer" in error for error in errors))

    def test_unknown_domain_is_reported(self):
        changed = deepcopy(self.project.use_cases)
        changed[0]["group"] = "UNKNOWN_DOMAIN"
        invalid_project = self.project.__class__(
            data_dir=self.project.data_dir,
            groups=self.project.groups,
            actors=self.project.actors,
            business_rules=self.project.business_rules,
            use_cases=changed,
        )
        errors = validate_project(invalid_project)
        self.assertTrue(any("unknown domain 'UNKNOWN_DOMAIN'" in error for error in errors))

    def test_use_case_in_wrong_domain_folder_is_reported(self):
        changed = deepcopy(self.project.use_cases)
        project_create = next(item for item in changed if item["key"] == "UC-PROJECT-CREATE")
        project_create["_source_file"] = "sprint-work/UC-PROJECT-CREATE.yml"
        invalid_project = self.project.__class__(
            data_dir=self.project.data_dir,
            groups=self.project.groups,
            actors=self.project.actors,
            business_rules=self.project.business_rules,
            use_cases=changed,
        )
        errors = validate_project(invalid_project)
        self.assertTrue(
            any(
                "file must be inside use_cases/project-membership/" in error
                for error in errors
            )
        )

    def test_string_alternative_flow_and_exception_can_be_generated(self):
        use_case = {
            "alternative_flows": ["The actor cancels the operation."],
            "exceptions": ["The database is unavailable."],
        }
        self.assertEqual(
            "AF-01: The actor cancels the operation.",
            _alternative_flow_text(use_case),
        )
        self.assertEqual(
            "EX-01: The database is unavailable.",
            _exception_text(use_case),
        )

    def test_docx_can_embed_diagram_from_manifest(self):
        one_pixel_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        root = ROOT / "tests" / "_tmp_docx_diagram"
        if root.exists():
            shutil.rmtree(root)
        root.mkdir()
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        try:
            diagram_dir = root / "diagrams"
            output_dir = root / "output"
            diagram_dir.mkdir()
            output_dir.mkdir()
            image_path = diagram_dir / "use-case-overview.png"
            image_path.write_bytes(one_pixel_png)
            (diagram_dir / "manifest.json").write_text(
                json.dumps(
                    {
                        "diagrams": [
                            {
                                "key": "overview",
                                "title": "Project Management System Use Case Model",
                                "files": {"png": image_path.name},
                            }
                        ]
                    }
                ),
                encoding="utf-8",
            )

            generated = _generate_docx(
                self.project,
                self.project.use_cases[:1],
                output_dir,
                diagram_dir,
            )
            document = Document(generated)

            self.assertEqual(1, len(document.inline_shapes))
            self.assertIn(
                "II.5.2.1 Use Case Diagram",
                [paragraph.text for paragraph in document.paragraphs],
            )
            self.assertIn(
                "II.5.2.2 Use Case Descriptions",
                [paragraph.text for paragraph in document.paragraphs],
            )
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_business_rule_outputs_include_all_rules(self):
        root = ROOT / "tests" / "_tmp_business_rules"
        if root.exists():
            shutil.rmtree(root)
        root.mkdir()
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        try:
            markdown_path = _generate_business_rules_markdown(self.project, root)
            docx_path = _generate_business_rules_docx(self.project, root)

            markdown = markdown_path.read_text(encoding="utf-8")
            document = Document(docx_path)
            table_rows = sum(len(table.rows) - 1 for table in document.tables)

            self.assertIn("# Business Rules", markdown)
            self.assertIn("BR-01", markdown)
            self.assertIn(f"BR-{len(self.project.business_rules):02d}", markdown)
            self.assertEqual(1, markdown.count("| ID | Business Rule | Description |"))
            self.assertEqual(1, len(document.tables))
            self.assertEqual(len(self.project.business_rules), table_rows)
            self.assertEqual("II.4 Business Rules", document.paragraphs[0].text)
            self.assertIsNone(
                document.tables[0].rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
            )
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_use_case_summary_table_has_balanced_columns_and_navy_header(self):
        root = ROOT / "tests" / "_tmp_summary_table_style"
        if root.exists():
            shutil.rmtree(root)
        root.mkdir()
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        try:
            generated = _generate_docx(
                self.project,
                self.project.use_cases,
                root,
                None,
            )
            document = Document(generated)
            summary_table = document.tables[0]

            self.assertIsNone(
                summary_table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
            )

            self.assertEqual(
                [1.6, 3.6, 3.7, 8.1],
                [round(column.width.cm, 1) for column in summary_table.columns],
            )
            for cell in summary_table.rows[0].cells:
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                self.assertIsNotNone(shading)
                self.assertEqual("1F4E78", shading.get(qn("w:fill")))
                visible_runs = [
                    run for paragraph in cell.paragraphs for run in paragraph.runs if run.text
                ]
                self.assertTrue(visible_runs)
                self.assertTrue(
                    all(run.font.color.rgb == RGBColor(255, 255, 255) for run in visible_runs)
                )
                self.assertTrue(all(run.font.size.pt == 12 for run in visible_runs))

            summary_body_runs = [
                run
                for paragraph in summary_table.rows[1].cells[1].paragraphs
                for run in paragraph.runs
                if run.text
            ]
            self.assertTrue(summary_body_runs)
            self.assertTrue(all(run.font.size.pt == 11 for run in summary_body_runs))

            detail_table = document.tables[1]
            self.assertTrue(
                all(
                    row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
                    for table in document.tables
                    for row in table.rows
                )
            )
            label_runs = [
                run
                for paragraph in detail_table.rows[0].cells[0].paragraphs
                for run in paragraph.runs
                if run.text
            ]
            value_runs = [
                run
                for paragraph in detail_table.rows[0].cells[1].paragraphs
                for run in paragraph.runs
                if run.text
            ]
            self.assertTrue(label_runs)
            self.assertTrue(value_runs)
            self.assertTrue(all(run.font.size.pt == 12 for run in label_runs))
            self.assertTrue(all(run.font.size.pt == 11 for run in value_runs))
            for row in detail_table.rows:
                shading = row.cells[0]._tc.get_or_add_tcPr().find(qn("w:shd"))
                self.assertIsNone(shading)
        finally:
            shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
